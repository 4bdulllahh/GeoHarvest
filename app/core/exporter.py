"""Excel and CSV export.

Only columns that have a value in at least one lead are written, so a quick-mode
export has no empty AI columns. Text from Google is untrusted: a business named
"=HYPERLINK(...)" must reach the sheet as text, never as a formula.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from app.core.places import FIELDS

LINK_COLUMNS = {"website", "maps_url", "reviews_url", "facebook", "instagram", "linkedin",
                "twitter", "youtube", "tiktok", "photo", "owner_linkedin"}
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # Excel refuses these outright
_PHONE = re.compile(r"^[+-][\d\s().-]+$")
EXCEL_CELL_LIMIT = 32767


def columns_in_use(leads: list[dict], keys: list[str] | None = None) -> list[tuple[str, str]]:
    """Columns with a value somewhere, limited to ``keys`` (the ones on screen) if given."""
    return [(key, header) for key, header in FIELDS
            if (keys is None or key in keys)
            and any(lead.get(key) not in (None, "", []) for lead in leads)]


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        value = ", ".join(map(str, value))
    return _CONTROL.sub("", str(value))[:EXCEL_CELL_LIMIT]


def safe_csv(value) -> str:
    """Stop spreadsheet apps treating a cell as a formula; phone numbers stay as they are."""
    text = _text(value)
    if text[:1] in ("=", "@") or (text[:1] in ("+", "-") and not _PHONE.match(text)):
        return "'" + text
    return text


def to_csv(leads: list[dict], path: Path, keys: list[str] | None = None) -> int:
    columns = columns_in_use(leads, keys)
    # utf-8-sig: the byte-order mark is what makes Excel read accents and Arabic correctly
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow([header for _key, header in columns])
        for lead in leads:
            writer.writerow([safe_csv(lead.get(key)) for key, _header in columns])
    return len(leads)


def to_xlsx(leads: list[dict], path: Path, keys: list[str] | None = None) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    columns = columns_in_use(leads, keys)
    book = Workbook()
    sheet = book.active
    sheet.title = "Leads"
    sheet.append([header for _key, header in columns])
    header_fill = PatternFill("solid", fgColor="2B5F96")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    link_font = Font(color="0563C1", underline="single")
    for row_number, lead in enumerate(leads, start=2):
        for col_number, (key, _header) in enumerate(columns, start=1):
            value = lead.get(key)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                sheet.cell(row_number, col_number, value)
                continue
            text = _text(value)
            cell = sheet.cell(row_number, col_number)
            cell.value = text
            cell.data_type = "s"  # never a formula, whatever the text starts with
            if key in LINK_COLUMNS and text.startswith(("http://", "https://")):
                cell.hyperlink = text
                cell.font = link_font

    for col_number, (key, header) in enumerate(columns, start=1):
        sample = [len(_text(lead.get(key))) for lead in leads[:200]]
        width = max([len(header)] + sample) + 2
        sheet.column_dimensions[get_column_letter(col_number)].width = min(max(width, 8), 50)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)
    return len(leads)


# CRM imports: each CRM's own column names for a company (account, organization), so its
# import screen matches them up by itself. "first_email" is the business's first email.
CRMS = {
    "hubspot": ("HubSpot", [
        ("Company name", "name"), ("Company domain name", "domain"), ("Website URL", "website"),
        ("Phone number", "phone"), ("Email", "first_email"), ("Street address", "street"),
        ("City", "city"), ("State/Region", "state"), ("Postal code", "postal_code"),
        ("Country/Region", "country_code"), ("Description", "description"),
        ("Contact name", "owner_name"), ("LinkedIn company page", "linkedin"),
        ("Facebook company page", "facebook"), ("Google Maps link", "maps_url")]),
    "pipedrive": ("Pipedrive", [
        ("Organization - Name", "name"), ("Organization - Address", "address"),
        ("Organization - Website", "website"), ("Person - Name", "owner_name"),
        ("Person - Email", "first_email"), ("Person - Phone", "phone"),
        ("Note - Content", "category")]),
    "zoho": ("Zoho CRM", [
        ("Account Name", "name"), ("Phone", "phone"), ("Website", "website"),
        ("Email", "first_email"), ("Industry", "category"), ("Billing Street", "street"),
        ("Billing City", "city"), ("Billing State", "state"), ("Billing Code", "postal_code"),
        ("Billing Country", "country_code"), ("Description", "description"),
        ("Contact Name", "owner_name"), ("Rating", "rating")]),
    "salesforce": ("Salesforce", [
        ("Account Name", "name"), ("Phone", "phone"), ("Website", "website"),
        ("Industry", "category"), ("Billing Street", "street"), ("Billing City", "city"),
        ("Billing State/Province", "state"), ("Billing Zip/Postal Code", "postal_code"),
        ("Billing Country", "country_code"), ("Description", "description"),
        ("Email", "first_email"), ("Contact Name", "owner_name")]),
}


def _crm_value(lead: dict, key: str):
    if key == "first_email":
        return (str(lead.get("emails") or "").split(",")[0]).strip()
    if key == "street" and not lead.get("street"):
        return lead.get("address")  # some countries' addresses have no street part
    return lead.get(key)


def to_crm(leads: list[dict], path: Path, crm: str) -> int:
    columns = CRMS[crm][1]
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow([header for header, _key in columns])
        for lead in leads:
            writer.writerow([safe_csv(_crm_value(lead, key)) for _header, key in columns])
    return len(leads)


def to_tsv(leads: list[dict], keys: list[str] | None = None) -> str:
    """Tab-separated text for pasting into Google Sheets (which reads it as cells)."""
    columns = columns_in_use(leads, keys)
    rows = [[header for _key, header in columns]]
    # Sheets reads a pasted "+971..." as a formula; a leading ' keeps it text and is not shown
    rows += [["'" + text if (text := " ".join(_text(lead.get(key)).split()))[:1] in "=+-@"
              else text for key, _header in columns] for lead in leads]
    return "\n".join("\t".join(row) for row in rows)


def export(leads: list[dict], path: Path, keys: list[str] | None = None) -> int:
    path = Path(path)
    return (to_csv if path.suffix.lower() == ".csv" else to_xlsx)(leads, path, keys)
